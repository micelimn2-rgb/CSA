"""Consulta de clientes con deuda y registro de acciones (CRM).

Se importa el Excel semanal "Consulta de Clientes con Deuda": cada fila es la deuda de un
cliente a una fecha (por tramos de días) con su vendedor, y las columnas que siguen a VENDEDOR
(MM, A. FALCON, JG...) son anotaciones de cada persona, que pasan al registro de acciones.
"""
import io
import json
import re
import unicodedata
from datetime import date, datetime

from fastapi import APIRouter, Body, File, HTTPException, UploadFile

from . import db

router = APIRouter(prefix="/api/clientes")
MAX_EXCEL = 10 * 1024 * 1024
TRAMO = re.compile(r"^\+?\s*\d+\s*(-\s*\d+)?\s*\+?$")  # "0-30", "90-120", "+120"


def clave(nombre: str) -> str:
    """Nombre normalizado para buscar y para no duplicar clientes ("Pérez  S.R.L " = "PEREZ S.R.L")."""
    sin_tildes = unicodedata.normalize("NFKD", nombre).encode("ascii", "ignore").decode()
    return " ".join(sin_tildes.upper().split())


def _ahora() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# ---------------------------------------------------------------- Lectura del Excel

def _filas_excel(contenido: bytes, nombre: str) -> list[list]:
    """Devuelve las filas de la primera hoja; las fechas como date."""
    if contenido[:4] == b"PK\x03\x04" or nombre.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        hoja = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True).worksheets[0]
        return [[c.date() if isinstance(c, datetime) else c for c in fila] for fila in hoja.iter_rows(values_only=True)]
    import xlrd
    libro = xlrd.open_workbook(file_contents=contenido)
    hoja = libro.sheet_by_index(0)
    filas = []
    for r in range(hoja.nrows):
        fila = []
        for c in range(hoja.ncols):
            celda = hoja.cell(r, c)
            if celda.ctype == xlrd.XL_CELL_DATE:
                fila.append(xlrd.xldate.xldate_as_datetime(celda.value, libro.datemode).date())
            else:
                fila.append(celda.value)
        filas.append(fila)
    return filas


def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def _numero(v) -> float:
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    t = _texto(v).replace("$", "").replace(" ", "")
    if not t:
        return 0.0
    if "," in t:  # formato argentino 1.234,56
        t = t.replace(".", "").replace(",", ".")
    try:
        return round(float(t), 2)
    except ValueError:
        return 0.0


def _fecha(v) -> str | None:
    if isinstance(v, date):
        return v.isoformat()
    t = _texto(v)
    m = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})$", t)
    if m:
        d, mes, a = map(int, m.groups())
        try:
            return date(a + 2000 if a < 100 else a, mes, d).isoformat()
        except ValueError:
            return None
    try:
        return date.fromisoformat(t[:10]).isoformat()
    except ValueError:
        return None


def leer_excel(contenido: bytes, nombre: str = "") -> list[dict]:
    """Interpreta la planilla. Devuelve [{nombre, fecha, vendedor, tramos, notas: [(autor, texto)]}]."""
    try:
        filas = _filas_excel(contenido, nombre)
    except Exception as exc:
        raise HTTPException(422, f"No se pudo leer el Excel: {exc}")

    # La fila de encabezados es la que tiene VENDEDOR
    for i, fila in enumerate(filas):
        cabeceras = [_texto(c).upper() for c in fila]
        if "VENDEDOR" in cabeceras:
            break
    else:
        raise HTTPException(422, "No encontré la fila de encabezados (con la columna VENDEDOR).")

    col_vendedor = cabeceras.index("VENDEDOR")
    col_fecha = cabeceras.index("FECHA") if "FECHA" in cabeceras else None
    tramos = {j: _texto(fila[j]) for j, c in enumerate(cabeceras) if TRAMO.match(c)}
    col_nombre = next((j for j, c in enumerate(cabeceras) if "CLIENTE" in c), None)
    if col_nombre is None:  # la primera columna que no es fecha, tramo ni vendedor
        col_nombre = next(j for j in range(len(cabeceras)) if j not in tramos and j not in (col_fecha, col_vendedor))
    notas = {j: _texto(filas[i][j]) for j in range(col_vendedor + 1, len(cabeceras)) if cabeceras[j]}

    registros = []
    for fila in filas[i + 1:]:
        fila = list(fila) + [None] * (len(cabeceras) - len(fila))
        nombre_cliente = _texto(fila[col_nombre])
        if not nombre_cliente:  # renglones vacíos o de totales
            continue
        fecha = _fecha(fila[col_fecha]) if col_fecha is not None else None
        registros.append({
            "nombre": nombre_cliente,
            "fecha": fecha or date.today().isoformat(),
            "vendedor": _texto(fila[col_vendedor]) or None,
            "tramos": {t: _numero(fila[j]) for j, t in tramos.items()},
            "notas": [(autor, _texto(fila[j])) for j, autor in notas.items() if _texto(fila[j])],
        })
    if not registros:
        raise HTTPException(422, "El Excel no tiene clientes.")
    return registros


# ---------------------------------------------------------------- Base de datos

def _cliente_id(conn, nombre: str, vendedor: str | None) -> tuple[int, bool]:
    k = clave(nombre)
    row = conn.execute("SELECT id FROM clientes WHERE clave = ?", (k,)).fetchone()
    if row:
        return row["id"], False
    cur = conn.execute("INSERT INTO clientes (nombre, clave, vendedor) VALUES (?, ?, ?)", (nombre, k, vendedor))
    return cur.lastrowid, True


def importar(conn, registros: list[dict]) -> dict:
    nuevos = cortes = notas = 0
    for r in sorted(registros, key=lambda r: r["fecha"]):
        cid, es_nuevo = _cliente_id(conn, r["nombre"], r["vendedor"])
        nuevos += es_nuevo
        conn.execute(
            "INSERT INTO cliente_cortes (cliente_id, fecha, tramos, total, vendedor) VALUES (?, ?, ?, ?, ?) "
            "ON CONFLICT (cliente_id, fecha) DO UPDATE SET tramos = excluded.tramos, total = excluded.total, "
            "vendedor = excluded.vendedor",
            (cid, r["fecha"], json.dumps(r["tramos"]), round(sum(r["tramos"].values()), 2), r["vendedor"]))
        cortes += 1
        # El vendedor del cliente es el del corte más reciente
        conn.execute("UPDATE clientes SET vendedor = (SELECT vendedor FROM cliente_cortes WHERE cliente_id = ? "
                     "ORDER BY fecha DESC LIMIT 1) WHERE id = ?", (cid, cid))
        for autor, texto in r["notas"]:
            # La misma nota suele repetirse semana a semana: se registra una sola vez
            if not conn.execute("SELECT 1 FROM cliente_comentarios WHERE cliente_id = ? AND autor = ? AND texto = ?",
                                (cid, autor, texto)).fetchone():
                conn.execute("INSERT INTO cliente_comentarios (cliente_id, fecha, autor, texto, origen) "
                             "VALUES (?, ?, ?, ?, 'excel')", (cid, r["fecha"], autor, texto))
                notas += 1
    return {"filas": len(registros), "clientes_nuevos": nuevos, "cortes": cortes, "comentarios_nuevos": notas}


def _ultimo_corte_general(conn) -> str | None:
    return conn.execute("SELECT MAX(fecha) AS f FROM cliente_cortes").fetchone()["f"]


def _resumen(row, ultimo_general: str | None) -> dict:
    d = dict(row)
    d.pop("clave", None)
    d["tramos"] = json.loads(d["tramos"]) if d.get("tramos") else {}
    # Si el cliente no aparece en el último Excel importado, probablemente ya no debe
    d["en_ultimo_corte"] = bool(d.get("fecha")) and d["fecha"] == ultimo_general
    return d


def _obtener(conn, cliente_id: int):
    row = conn.execute("SELECT * FROM clientes WHERE id = ?", (cliente_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Cliente no encontrado")
    return row


# ---------------------------------------------------------------- API

CONSULTA = """
SELECT c.*, k.fecha, k.tramos, k.total,
       (SELECT COUNT(*) FROM cliente_comentarios m WHERE m.cliente_id = c.id) AS comentarios,
       (SELECT MAX(fecha) FROM cliente_comentarios m WHERE m.cliente_id = c.id) AS ultimo_comentario
FROM clientes c
LEFT JOIN cliente_cortes k ON k.cliente_id = c.id
     AND k.fecha = (SELECT MAX(fecha) FROM cliente_cortes WHERE cliente_id = c.id)
"""


@router.get("")
def buscar(q: str = "", limite: int = 50):
    palabras = clave(q).split()
    where = " AND ".join("c.clave LIKE ?" for _ in palabras) or "1"
    with db.connect() as conn:
        rows = conn.execute(f"{CONSULTA} WHERE {where} ORDER BY k.total DESC, c.nombre LIMIT ?",
                            [f"%{p}%" for p in palabras] + [max(1, min(limite, 500))]).fetchall()
        ultimo = _ultimo_corte_general(conn)
        total = conn.execute("SELECT COUNT(*) AS n FROM clientes").fetchone()["n"]
    return {"clientes": [_resumen(r, ultimo) for r in rows], "total_clientes": total, "ultimo_corte": ultimo}


@router.post("/importar")
async def importar_excel(archivo: UploadFile = File(...)):
    contenido = await archivo.read()
    if not contenido:
        raise HTTPException(422, "El archivo está vacío")
    if len(contenido) > MAX_EXCEL:
        raise HTTPException(413, "El archivo supera los 10 MB")
    registros = leer_excel(contenido, archivo.filename or "")
    with db.connect() as conn:
        resultado = importar(conn, registros)
        resultado["ultimo_corte"] = _ultimo_corte_general(conn)
    return resultado


@router.get("/{cliente_id}")
def ficha(cliente_id: int):
    with db.connect() as conn:
        _obtener(conn, cliente_id)
        cliente = _resumen(conn.execute(f"{CONSULTA} WHERE c.id = ?", (cliente_id,)).fetchone(),
                           _ultimo_corte_general(conn))
        cliente["cortes"] = [{**dict(r), "tramos": json.loads(r["tramos"])} for r in conn.execute(
            "SELECT fecha, tramos, total, vendedor FROM cliente_cortes WHERE cliente_id = ? ORDER BY fecha DESC",
            (cliente_id,))]
        cliente["historial"] = [dict(r) for r in conn.execute(
            "SELECT id, fecha, autor, texto, origen FROM cliente_comentarios WHERE cliente_id = ? "
            "ORDER BY fecha DESC, id DESC", (cliente_id,))]
    return cliente


@router.post("/{cliente_id}/comentarios", status_code=201)
def comentar(cliente_id: int, cuerpo: dict = Body(...)):
    autor = " ".join(str(cuerpo.get("autor") or "").split())[:60]
    texto = str(cuerpo.get("texto") or "").strip()[:4000]
    if not autor:
        raise HTTPException(422, "Poné tu nombre para que se sepa quién hizo la acción")
    if not texto:
        raise HTTPException(422, "El comentario está vacío")
    with db.connect() as conn:
        _obtener(conn, cliente_id)
        cur = conn.execute("INSERT INTO cliente_comentarios (cliente_id, fecha, autor, texto, origen) "
                           "VALUES (?, ?, ?, ?, 'web')", (cliente_id, _ahora(), autor, texto))
        return dict(conn.execute("SELECT id, fecha, autor, texto, origen FROM cliente_comentarios WHERE id = ?",
                                 (cur.lastrowid,)).fetchone())
